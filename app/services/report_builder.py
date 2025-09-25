# app/services/report_builder.py
import os
from datetime import datetime, date
from typing import Dict, Any, List

from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from PIL import Image

from core.config import OUTPUT_DIR


TITLE_FONT = "Times New Roman"
BODY_FONT = "Times New Roman"


def _ensure_reports_dir() -> str:
    reports_dir = os.path.join(OUTPUT_DIR, "reports")
    os.makedirs(reports_dir, exist_ok=True)
    return reports_dir


def _set_doc_defaults(doc: Document):
    # Márgenes 2.54cm (estándar)
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

    # Fuente por defecto
    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(12)

    # Espaciado 1.5 aprox
    style.paragraph_format.line_spacing = 1.5


def _available_width_inches(doc: Document) -> float:
    s = doc.sections[0]
    avail = s.page_width - s.left_margin - s.right_margin
    # EMU -> inches
    return float(avail) / 914400.0


def _add_title(doc: Document, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.name = TITLE_FONT
    run.font.size = Pt(20)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_subtitle(doc: Document, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.name = TITLE_FONT
    run.font.size = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def _format_money(val: float) -> str:
    try:
        return "${:,.2f}".format(float(val))
    except Exception:
        return str(val)


def _pick_images(payload: Dict[str, Any]) -> List[str]:
    paths = []
    for k in ("screenshot_path", "screenshot_historial_path"):
        p = payload.get(k)
        if p and os.path.exists(p):
            paths.append(p)
    return paths


def _human_name(tipo: str) -> str:
    mapping = {
        "ruc": "SRI – RUC",
        "deudas": "SRI – Deudas Firmes/Impugnadas",
        "denuncias": "Fiscalía – Denuncias",
        "mercado_valores": "Superintendencia – Mercado de Valores",
        "interpol": "INTERPOL – Notificaciones",
        "google": "Google – Búsqueda",
        "contraloria": "Contraloría – DDJJ",
        "supercias_persona": "Superintendencia – Consulta de Persona",
        "predio_quito": "GAD Quito – Predios",
        "predio_manta": "GAD Manta – Predios",
        "funcion_judicial": "Función Judicial – Procesos Judiciales",
    }
    return mapping.get(tipo, tipo)


def build_report_docx(job_id: str, meta: Dict[str, Any], results: Dict[str, Any]) -> str:
    """
    Construye un DOCX profesional (APA-like) con portada, secciones por consulta y conclusión.
    Retorna la ruta absoluta del archivo .docx generado.
    """
    reports_dir = _ensure_reports_dir()

    # Documento
    doc = Document()
    _set_doc_defaults(doc)

    # Portada / Encabezado
    _add_title(doc, "Revisión de Función Judicial")
    #tipo_alerta = str(meta.get("tipo_alerta", "General"))
    #monto = meta.get("monto_usd", None)
    #fecha_alerta = meta.get("fecha_alerta")

    # Normalizar fecha
    if isinstance(fecha_alerta, str):
        try:
            fecha_alerta = date.fromisoformat(fecha_alerta)
        except Exception:
            fecha_alerta = None

    doc.add_paragraph(f"Tipo de alerta: {tipo_alerta}")
    if monto is not None:
        doc.add_paragraph(f"Monto (USD): {_format_money(monto)}")
    if fecha_alerta:
        doc.add_paragraph(f"Fecha de la alerta: {fecha_alerta.isoformat()}")
    #doc.add_paragraph(f"Job ID: {job_id}")
    doc.add_paragraph(f"Fecha de generación: {datetime.now().isoformat(sep=' ', timespec='seconds')}")

    doc.add_paragraph("")  # espacio

    # Secciones por consulta
    width_in = _available_width_inches(doc)
    max_w = max(3.0, width_in)  # seguridad

    figura_idx = 1
    for tipo, payload in results.items():
        _add_subtitle(doc, f"Consulta: {_human_name(tipo)}")

        # Breve comentario/escenario (2 líneas máximo)
        scenario = payload.get("scenario")
        if scenario:
            doc.add_paragraph(f"Resumen: {scenario}")
        else:
            doc.add_paragraph("Se adjunta evidencia visual de la consulta realizada.")

        # Insertar imágenes si existen
        imgs = _pick_images(payload)
        if not imgs:
            doc.add_paragraph("No se generaron capturas para esta consulta.")
        else:
            for img_path in imgs:
                try:
                    # Ajuste proporcional al ancho disponible
                    with Image.open(img_path) as im:
                        doc.add_picture(img_path, width=Inches(max_w))
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap.add_run(f"Figura {figura_idx}. Evidencia – {_human_name(tipo)}").italic = True
                    figura_idx += 1
                except Exception:
                    doc.add_paragraph(f"[Aviso] Falló al insertar la imagen: {img_path}")

        doc.add_paragraph("")  # separación

    # Conclusión (placeholder – sin LLM por ahora)
    _add_subtitle(doc, "Conclusión")
    concl = (
        "Con base en las evidencias adjuntas, se confirma que las consultas fueron ejecutadas en las "
        "fuentes oficiales indicadas. Este informe no realiza extracción automática de montos; por lo "
        "que la validación del valor transaccionado debe contrastarse visualmente con las capturas. "
        "De ser requerido, en una siguiente fase se integrará análisis asistido por LLM con pautas "
        "para relacionar hallazgos con el monto reportado."
    )
    doc.add_paragraph(concl)

    # Guardar
    filename = f"report_{job_id}.docx"
    out_path = os.path.abspath(os.path.join(reports_dir, filename))
    doc.save(out_path)
    return out_path
